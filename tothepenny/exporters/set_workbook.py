# SPDX-License-Identifier: AGPL-3.0-only OR LicenseRef-Commercial
# Copyright (C) 2026 Paul Wood FRSA. tothepenny by Paul Wood FRSA; see NOTICE and COMMERCIAL.md.

"""One workbook for a set of statements: every transaction from a reconciled statement, every statement and whether it reconciled,
where an account's statements do not join up, and the transfers matched between accounts in the set."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .multi_account_excel import MONEY, header, literal, style, table


BANK_NAMES = {'hsbc': 'HSBC', 'tsb': 'TSB', 'lcl': 'LCL', 'natwest': 'NatWest', 'credit_agricole': 'Crédit Agricole',
              'pagseguro': 'PagSeguro', 'metro': 'Metro Bank'}


def bank_name(bank) -> str:
    return BANK_NAMES.get((bank or '').lower(), (bank or '').replace('_', ' ').title())


def _money(value):
    if value in (None, ''):
        return None
    try:
        return round(float(value), 2)
    except (TypeError, ValueError):
        return None


def _day(value):
    try:
        return date.fromisoformat(str(value)[:10]) if value else None
    except ValueError:
        return None


def _sheet(wb, title, labels, rows, money_cols=(), date_cols=(), name=None, widths=None, empty=None):
    ws = wb.create_sheet(title)
    header(ws, 1, labels)
    if not rows and empty:
        ws.cell(2, 1, empty)
    for r, values in enumerate(rows, start=2):
        for c, value in enumerate(values, start=1):
            cell = ws.cell(r, c)
            if c in money_cols:
                cell.value = _money(value)
                cell.number_format = MONEY
            elif c in date_cols:
                cell.value = _day(value)
                cell.number_format = 'DD/MM/YYYY'
            else:
                literal(cell, '' if value is None else value)
    style(ws)
    if rows:
        table(ws, len(rows) + 1, len(labels), name or title.replace(' ', ''))
    for col, width in (widths or {}).items():
        ws.column_dimensions[get_column_letter(col)].width = width
    return ws


def write_set_workbook(summary, rows, pairs, findings, path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    _sheet(wb, 'Transactions',
           ['Date', 'Account', 'Statement', 'Page', 'Description', 'Paid in', 'Paid out', 'Balance', 'Matched transfer'],
           [[r['Date'], r['Account'], r['Source file'], r['Source page'], r['Description'], r['Paid In'],
             r['Withdrawn'], r['Balance'], r.get('Matched transfer', '')] for r in rows],
           money_cols=(6, 7, 8), date_cols=(1,), widths={3: 28, 5: 48, 9: 36},
           empty='No statement reconciled with the bank\'s figures, so there are no transactions to list.')

    statements = []
    for result in summary.results:
        if result.skipped:
            reconciled, why = 'No', 'Scanned: no text to read'
        elif result.reconciled:
            reconciled, why = 'Yes', ''
        else:
            reconciled = 'No'
            why = result.error or (result.warnings[-1] if result.warnings else 'Did not reconcile to the bank\'s figures')
        statements.append([result.file, bank_name(result.bank), result.account, result.period_start,
                           result.period_end, result.opening, result.closing, result.transactions, reconciled, why])
    _sheet(wb, 'Statements',
           ['Statement', 'Bank', 'Account', 'From', 'To', 'Opening', 'Closing', 'Transactions', 'Reconciled', 'Why not'],
           statements, money_cols=(6, 7), date_cols=(4, 5), widths={1: 32, 10: 60})

    gaps = [f for f in findings if f['kind'] != 'joined']
    _sheet(wb, 'Gaps',
           ['Account', 'What', 'Statement', 'Ends', 'Next statement', 'Starts', 'Closing', 'Next opening',
            'Difference'],
           [[f['account'], f['kind'].capitalize(), f['from_file'], f['from_end'], f['to_file'], f['to_start'],
             f['closing'], f['next_opening'], f['unexplained']] for f in gaps],
           money_cols=(7, 8, 9), date_cols=(4, 6), widths={3: 28, 5: 28},
           empty='Every reconciled statement joins the next for its account.' if findings
           else 'No account has two reconciled statements to join.')

    _sheet(wb, 'Transfers',
           ['Date out', 'From account', 'Date in', 'To account', 'Amount', 'Why they are the same money'],
           [[p['out']['Date'], p['out']['Account'], p['in']['Date'], p['in']['Account'], p['out']['Withdrawn'],
             p['evidence']] for p in pairs],
           money_cols=(5,), date_cols=(1, 3), widths={6: 40},
           empty='No transfers between accounts in this set were matched.')

    for ws in wb.worksheets:
        ws.cell(1, 1).font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
