# SPDX-License-Identifier: AGPL-3.0-only OR LicenseRef-Commercial
# Copyright (C) 2026 Paul Wood FRSA. tothepenny by Paul Wood FRSA; see NOTICE and COMMERCIAL.md.

"""Account-aware Excel output for statements with independent balance streams."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


BLUE = '294B6B'
PALE = 'EEF3F7'
MONEY = '#,##0.00;[Red](#,##0.00);–'


def literal(cell, value):
    """Keep source strings literal, including strings that begin with '='."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = 's'


def header(ws, row, labels):
    for col, label in enumerate(labels, 1):
        cell = ws.cell(row, col, label)
        cell.fill = PatternFill('solid', fgColor=BLUE)
        cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = 30


def table(ws, end_row, end_col, name):
    if end_row < 2:
        return
    obj = Table(displayName=name, ref=f'A1:{get_column_letter(end_col)}{end_row}')
    obj.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(obj)
    ws.freeze_panes = 'C2'


def style(ws):
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows():
        for cell in row:
            if not cell.font.bold:
                cell.font = Font(name='Arial', size=10)
            cell.alignment = Alignment(vertical='center', wrap_text=cell.column in (3, 4))
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 17


def export_multi_account(result, output_path: Path):
    wb = Workbook()
    summary = wb.active
    summary.title = 'Summary'
    monthly = wb.create_sheet('Monthly totals')
    tx = wb.create_sheet('Transactions')
    reverted = wb.create_sheet('Reverted entries')
    audit = wb.create_sheet('Extraction Log')
    accounts = result.account_summaries
    end = max(2, len(result.transactions) + 1)
    account_rows = {a['account_name']: i + 8 for i, a in enumerate(accounts)}

    summary['A2'] = 'Revolut statement extraction'
    summary['A2'].font = Font(name='Arial', size=14, bold=True)
    literal(summary['A3'], result.statement.account_holder or 'Account holder not stated')
    summary['A4'] = f"{result.statement.statement_start_date:%d %B %Y} to {result.statement.statement_end_date:%d %B %Y} (GBP)"
    summary['A6'] = 'Printed balance summary'
    header(summary, 7, ['Account', 'Opening balance', 'Money out', 'Money in', 'Closing balance'])
    for r, account in enumerate(accounts, 8):
        literal(summary.cell(r, 1), account['account_name'])
        for c, key in enumerate(('opening_balance', 'money_out', 'money_in', 'closing_balance'), 2):
            summary.cell(r, c, account[key]).number_format = MONEY
    source_end = 7 + len(accounts)
    summary.cell(source_end + 1, 1, 'Total')
    for c in range(2, 6):
        letter = get_column_letter(c)
        summary.cell(source_end + 1, c, f'=SUM({letter}8:{letter}{source_end})').number_format = MONEY

    check_title = source_end + 4
    summary.cell(check_title, 1, 'Extraction reconciliation')
    check_header = check_title + 1
    header(summary, check_header, ['Account', 'Completed rows', 'Money out', 'Money in', 'Out difference', 'In difference', 'Closing difference', 'Row balance issues'])
    for i, account in enumerate(accounts):
        r = check_header + i + 1
        src = account_rows[account['account_name']]
        literal(summary.cell(r, 1), account['account_name'])
        summary.cell(r, 2, f'=COUNTIFS(Transactions!$B$2:$B${end},A{r})')
        summary.cell(r, 3, f'=SUMIFS(Transactions!$E$2:$E${end},Transactions!$B$2:$B${end},A{r})')
        summary.cell(r, 4, f'=SUMIFS(Transactions!$F$2:$F${end},Transactions!$B$2:$B${end},A{r})')
        summary.cell(r, 5, f'=ROUND(C{r}-C{src},2)')
        summary.cell(r, 6, f'=ROUND(D{r}-D{src},2)')
        summary.cell(r, 7, f'=ROUND(B{src}+D{r}-C{r}-E{src},2)')
        summary.cell(r, 8, f'=COUNTIFS(Transactions!$B$2:$B${end},A{r},Transactions!$I$2:$I${end},"<>0")')
        for c in range(3, 8):
            summary.cell(r, c).number_format = MONEY
    checks_end = check_header + len(accounts)
    summary.conditional_formatting.add(f'E{check_header+1}:H{checks_end}', CellIsRule(operator='notEqual', formula=['0'], fill=PatternFill('solid', fgColor='FDE0DD')))
    notes_row = checks_end + 3
    for offset, text in enumerate([
        'Money in and out include transfers between the accounts shown. They are not income and expenditure totals.',
        f'{len(result.excluded_transactions)} reverted entries are retained separately and excluded from completed-transaction totals.',
        'A zero difference means the extracted amount matches the printed source control.',
        f'Source: {result.source_file}, balance summary on page 1.',
    ]):
        literal(summary.cell(notes_row + offset, 1), text)
        summary.merge_cells(start_row=notes_row + offset, start_column=1, end_row=notes_row + offset, end_column=8)
        summary.cell(notes_row + offset, 1).alignment = Alignment(wrap_text=True, vertical='center')
        summary.row_dimensions[notes_row + offset].height = 30

    labels = ['Date', 'Account', 'Description', 'Details', 'Money out', 'Money in', 'Balance', 'Source page', 'Balance difference']
    header(tx, 1, labels)
    previous_rows = {}
    for r, txn in enumerate(result.transactions, 2):
        values = [txn.date, txn.account_name, txn.description, txn.details or '', txn.money_out, txn.money_in, txn.balance, txn.page_number]
        for c, value in enumerate(values, 1):
            literal(tx.cell(r, c), value)
        tx.cell(r, 1).number_format = 'dd/mm/yyyy'
        for c in (5, 6, 7, 9):
            tx.cell(r, c).number_format = MONEY
        opening = f'G{previous_rows[txn.account_name]}' if txn.account_name in previous_rows else f'Summary!B{account_rows[txn.account_name]}'
        tx.cell(r, 9, f'=ROUND(G{r}-({opening}+F{r}-E{r}),2)')
        tx.row_dimensions[r].height = 42 if len(txn.details or '') > 65 else 30
        previous_rows[txn.account_name] = r
    table(tx, end, len(labels), 'CompletedTransactions')
    tx.conditional_formatting.add(f'I2:I{end}', CellIsRule(operator='notEqual', formula=['0'], fill=PatternFill('solid', fgColor='FDE0DD')))

    header(reverted, 1, ['Start date', 'Account', 'Description', 'Details', 'Money out', 'Money in', 'Source page', 'Status'])
    for r, txn in enumerate(result.excluded_transactions, 2):
        for c, value in enumerate([txn.date, txn.account_name, txn.description, txn.details or '', txn.money_out, txn.money_in, txn.page_number, txn.status], 1):
            literal(reverted.cell(r, c), value)
        reverted.cell(r, 1).number_format = 'dd/mm/yyyy'
        for c in (5, 6):
            reverted.cell(r, c).number_format = MONEY
        reverted.row_dimensions[r].height = 42
    table(reverted, len(result.excluded_transactions) + 1, 8, 'RevertedEntries')

    header(monthly, 1, ['Month', 'Account', 'Completed rows', 'Money out', 'Money in', 'Net movement'])
    start = result.statement.statement_start_date.replace(day=1)
    finish = result.statement.statement_end_date
    r = 2
    while start <= finish:
        next_month = start.replace(year=start.year+1, month=1) if start.month == 12 else start.replace(month=start.month+1)
        for account in accounts:
            monthly.cell(r, 1, start).number_format = 'mmm yyyy'
            literal(monthly.cell(r, 2), account['account_name'])
            criteria = f'Transactions!$B$2:$B${end},B{r},Transactions!$A$2:$A${end},">="&A{r},Transactions!$A$2:$A${end},"<"&DATE({next_month.year},{next_month.month},1)'
            monthly.cell(r, 3, f'=COUNTIFS({criteria})')
            monthly.cell(r, 4, f'=SUMIFS(Transactions!$E$2:$E${end},{criteria})').number_format = MONEY
            monthly.cell(r, 5, f'=SUMIFS(Transactions!$F$2:$F${end},{criteria})').number_format = MONEY
            monthly.cell(r, 6, f'=E{r}-D{r}').number_format = MONEY
            r += 1
        start = next_month
    table(monthly, r - 1, 6, 'MonthlyAccountTotals')

    audit['A2'] = 'Extraction record'
    audit['A2'].font = Font(name='Arial', size=14, bold=True)
    for r, (key, value) in enumerate([
        ('Method', result.extraction_method), ('Cloud extraction APIs used', 'No'),
        ('Source file', result.source_file), ('Source SHA-256', result.source_sha256),
        ('PDF pages', len(result.page_coverage)), ('Completed transactions', result.transaction_count),
        ('Reverted entries', len(result.excluded_transactions)), ('Unresolved issues', len(result.warnings)),
    ], 4):
        audit.cell(r, 1, key)
        literal(audit.cell(r, 2), value)
    header(audit, 14, ['Source page', 'Date rows found', 'Completed rows', 'Reverted rows', 'Unaccounted rows'])
    for r, page in enumerate(result.page_coverage, 15):
        for c, key in enumerate(('page', 'date_anchors', 'completed_rows', 'reverted_rows'), 1):
            audit.cell(r, c, page[key])
        audit.cell(r, 5, f'=B{r}-C{r}-D{r}')
    audit.freeze_panes = 'A15'
    for ws in wb:
        style(ws)
    summary.column_dimensions['A'].width = 32
    for ws in (tx, reverted):
        ws.column_dimensions['A'].width = 13
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 42
        ws.column_dimensions['D'].width = 68
    monthly.column_dimensions['B'].width = 32
    for col in ('C', 'D', 'E', 'F'):
        monthly.column_dimensions[col].width = 20
    audit.column_dimensions['A'].width = 28
    audit.column_dimensions['B'].width = 90
    audit.column_dimensions['C'].width = 19
    audit.column_dimensions['D'].width = 18
    audit.column_dimensions['E'].width = 22
    # Restore explanatory wrapping after the shared style pass.
    for r in range(notes_row, notes_row + 4):
        summary.cell(r, 1).alignment = Alignment(wrap_text=True, vertical='center')
    wb.calculation.fullCalcOnLoad = True
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
