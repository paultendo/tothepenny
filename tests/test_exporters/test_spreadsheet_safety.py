# SPDX-License-Identifier: AGPL-3.0-only OR LicenseRef-Commercial
# Copyright (C) 2026 Paul Wood FRSA. tothepenny by Paul Wood FRSA; see NOTICE and COMMERCIAL.md.

"""Statement text must stay inert in exported spreadsheets. All data is synthetic."""
from openpyxl import Workbook, load_workbook

from tothepenny.utils.spreadsheet_safety import csv_safe, neutralise_formulas


def test_formula_text_from_a_statement_is_stored_as_text(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = '=HYPERLINK("http://example.invalid","click")'
    ws["A2"] = "TESCO STORES"
    assert neutralise_formulas(wb) == 1
    path = tmp_path / "out.xlsx"
    wb.save(path)
    cell = load_workbook(path).active["A1"]
    assert cell.data_type == "s"
    assert cell.value.startswith("=HYPERLINK")


def test_csv_fields_that_would_run_as_formulas_are_prefixed():
    assert csv_safe("=1+1") == "'=1+1"
    assert csv_safe("+44 20 0000 0000") == "'+44 20 0000 0000"
    assert csv_safe("-10.00") == "'-10.00"
    assert csv_safe("@SUM(A1)") == "'@SUM(A1)"
    assert csv_safe("statement.pdf") == "statement.pdf"
    assert csv_safe(12) == 12
