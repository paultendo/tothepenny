# SPDX-License-Identifier: AGPL-3.0-only OR LicenseRef-Commercial
# Copyright (C) 2026 Paul Wood FRSA. statement-reconciler by Paul Wood FRSA; see NOTICE and COMMERCIAL.md.

"""Synthetic native-PDF regression gates; no client statement data."""
import csv
import hashlib

from reportlab.pdfgen import canvas


class _Page:
    """Text placed at PDF points from the top-left, as a statement generator would place it."""

    def __init__(self, width, height):
        self.width, self.height, self.items = width, height, []

    def insert_text(self, point, text, fontsize=8):
        self.items.append((point[0], point[1], text, fontsize))


class _Document:
    """A minimal synthetic-PDF writer on reportlab (BSD), used only to build test statements."""

    def __init__(self):
        self.pages = []

    def new_page(self, width, height):
        page = _Page(width, height)
        self.pages.append(page)
        return page

    def save(self, path):
        pdf = canvas.Canvas(str(path))
        for page in self.pages:
            pdf.setPageSize((page.width, page.height))
            for x, y, text, size in page.items:
                pdf.setFont("Helvetica", size)
                pdf.drawString(x, page.height - y, text)
            pdf.showPage()
        pdf.save()

    def close(self):
        pass
import pytest
from openpyxl import load_workbook

from statement_reconciler.config.bank_config_loader import BankConfigLoader
from statement_reconciler.parsers.revolut_parser import RevolutParser
from statement_reconciler.pipeline import ExtractionPipeline


def statement_pdf(tmp_path, *, shift=0, surname='Example', fault=None):
    doc = _Document()
    page = doc.new_page(width=680, height=880)
    def put(x, y, text):
        page.insert_text((x + shift, y + shift), text, fontsize=8)
    def headings(y, reverted=False):
        for x, text in [(30, 'Start date' if reverted else 'Date'), (140, 'Description'), (365, 'Money out'), (455, 'Money in')]:
            put(x, y, text)
        if not reverted:
            put(560, y, 'Balance')
    def txn(y, desc, out=None, incoming=None, balance=None, date='2 Jan 2024'):
        put(30, y, date)
        put(140, y, desc)
        for x, value in ((365, out), (455, incoming), (560, balance)):
            if value is not None:
                put(x, y, f'£{value:.2f}')
    period = 'from 1 Jan 2024 to 31 Jan 2024'
    put(30, 25, 'GBP Statement' if fault != 'currency' else 'EUR Statement')
    put(30, 45, 'Revolut Ltd')
    put(30, 65, f'JANE {surname.upper()}')
    put(300, 85, 'IBAN')
    put(300, 100, 'Account Number'); put(430, 100, '12345678')
    put(300, 115, 'Sort Code'); put(430, 115, '04-29-09')
    put(30, 140, 'Balance summary')
    for x, text in [(30, 'Product'), (255, 'Opening balance'), (365, 'Money out'), (455, 'Money in'), (560, 'Closing balance')]:
        put(x, 160, text)
    for y, label, values in [(180, 'Account (e-money)', (10, 13, 20, 17)), (210, 'Personal and Group Pockets', (1, 0, 5, 6)), (240, 'Account (e-money)', (2, 1, 0, 1)), (275, 'Total', (13, 14, 25, 24 if fault != 'summary' else 25))]:
        put(30, y, label)
        for x, value in zip((255, 365, 455, 560), values): put(x, y, f'£{value:.2f}')
    put(30, 253, f'Alex {surname}')
    put(30, 310, 'Account transactions ' + period)
    headings(330)
    txn(355, '=merchant', incoming=20, balance=30)
    txn(385, 'Cash withdrawal', out=13, balance=17 if fault != 'balance' else 18,
        date='2 ??? 2024' if fault == 'date' else '2 Jan 2024')
    put(140, 398, 'Fee: £1.00'); put(365, 398, '£12.00' if fault != 'fee' else '£11.00')
    put(140, 410, 'Reference: terminal receipt')
    # Footer's small text starts fractionally above the left-hand card notice.
    put(190, 800, 'Your Retail current account is an e-money account')
    put(80, 800.3, 'Report lost or stolen card')
    page = doc.new_page(width=680, height=880)
    put(30, 40, 'Reverted ' + period)
    headings(60, True); txn(80, 'Cancelled purchase', out=7)
    put(30, 115, 'Personal and Group Pockets transactions ' + period)
    headings(135); txn(155, 'Pocket transfer', incoming=5, balance=6)
    put(30, 190, ("Unknown's account" if fault == 'account' else "Alex's account") + ' transactions ' + period)
    headings(210); txn(230, 'Purchase', out=1, balance=1)
    if fault == 'unassigned': put(365, 250, '£9.00')
    put(80, 800, 'Report lost or stolen card')
    path = tmp_path / f'synthetic-{shift}.pdf'
    doc.save(path); doc.close()
    return path


@pytest.mark.parametrize('shift,surname', [(0, 'Example'), (19, 'Sample')])
def test_independent_account_balances_fees_reverted_and_page_coverage(tmp_path, shift, surname):
    path = statement_pdf(tmp_path, shift=shift, surname=surname)
    result = RevolutParser().parse_pdf(path)
    assert result.success and result.balance_reconciled
    assert len(result.transactions) == 4
    assert [a['transaction_count'] for a in result.account_summaries] == [2, 1, 1]
    assert [a['calculated_closing'] for a in result.account_summaries] == [17, 6, 1]
    assert all(a['reconciled'] for a in result.account_summaries)
    assert result.statement.opening_balance == 13
    assert result.statement.closing_balance == 24
    assert result.statement.account_holder == f'JANE {surname.upper()}'
    assert result.transactions[1].money_out == 13
    assert 'Fee: £1.00' in result.transactions[1].details
    assert 'Amount before fee: £12.00' in result.transactions[1].details
    assert not any('Your Retail' in t.details for t in result.transactions)
    assert len(result.excluded_transactions) == 1
    assert result.excluded_transactions[0].status == 'Reverted'
    assert result.excluded_transactions[0].balance is None
    assert result.page_coverage == [
        dict(page=1, date_anchors=2, completed_rows=2, reverted_rows=0),
        dict(page=2, date_anchors=3, completed_rows=2, reverted_rows=1)]
    assert result.source_sha256 == hashlib.sha256(path.read_bytes()).hexdigest()


@pytest.mark.parametrize('fault,reason', [('currency', 'GBP'), ('summary', 'summary'), ('fee', 'Fee breakdown'), ('date', 'Unrecognized transaction date'), ('account', 'uniquely match'), ('unassigned', 'Unexplained amount')])
def test_malformed_statement_fails_closed(tmp_path, fault, reason):
    with pytest.raises(ValueError, match=reason):
        RevolutParser().parse_pdf(statement_pdf(tmp_path, fault=fault))


def test_bad_balance_returns_failed_result_and_pipeline_does_not_export(tmp_path):
    path = statement_pdf(tmp_path, fault='balance')
    output = tmp_path / 'bad.xlsx'
    result = ExtractionPipeline().process(path, output_path=output, perform_validation=False)
    assert not result.success
    assert not result.balance_reconciled
    assert not output.exists()
    assert result.account_summaries[0]['balance_mismatches'][0]['difference'] == 1


def test_bank_detection_does_not_guess_all_04_sort_codes_are_monzo():
    loader = BankConfigLoader()
    assert loader.detect_bank('Sort Code 04-29-09') is None
    assert loader.detect_bank('Revolut Ltd\nSort Code 04-29-09').bank_name == 'revolut'
    assert loader.detect_bank('Sort Code 04-00-04').bank_name == 'monzo'


def test_standard_pipeline_excel_and_csv_preserve_account_provenance(tmp_path):
    path = statement_pdf(tmp_path)
    xlsx = tmp_path / 'statement.xlsx'
    result = ExtractionPipeline().process(path, output_path=xlsx)
    assert result.success
    wb = load_workbook(xlsx)
    assert wb.sheetnames == ['Summary', 'Monthly totals', 'Transactions', 'Reverted entries', 'Extraction Log']
    assert wb['Transactions'].max_row == 5
    assert wb['Transactions']['C2'].value == '=merchant'
    assert wb['Transactions']['C2'].data_type == 's'
    assert wb['Transactions']['H5'].value == 2
    assert wb['Transactions']['I4'].value == '=ROUND(G4-(Summary!B9+F4-E4),2)'
    assert wb['Reverted entries']['H2'].value == 'Reverted'
    assert len(result.to_dict()['excluded_transactions']) == 1
    csv_path = tmp_path / 'statement.csv'
    csv_result = ExtractionPipeline().process(path, output_path=csv_path, export_format='csv')
    assert csv_result.success
    with csv_path.open(encoding='utf-8-sig') as f:
        rows = list(csv.DictReader(f))
    assert len(rows) == 4
    assert rows[-1]['Account'] == 'Alex Example'
    assert rows[-1]['Source Page'] == '2'
